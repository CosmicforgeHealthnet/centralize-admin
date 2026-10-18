"""Build the canonical Marketing Intelligence Phase 2 architecture DOCX.

The Markdown file remains the canonical editable source. This script creates a
review-ready Word companion with an editorial cover, document control, a TOC,
native tables, rendered diagram equivalents, running headers, and page numbers.
"""

from __future__ import annotations

import argparse
import math
import re
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_SOURCE = ROOT / "docs" / "CosmicForge_Marketing_Intelligence_Phase2_Complete_Architecture.md"
DEFAULT_OUTPUT = ROOT / "docs" / "CosmicForge_Marketing_Intelligence_Phase2_Complete_Architecture.docx"
DEFAULT_DIAGRAM_DIR = ROOT / "artifacts" / "phase2-architecture-diagrams"

NAVY = "16324F"
BLUE = "2E74B5"
TEAL = "1F7A8C"
INK = "243447"
MUTED = "5E6C7B"
PALE = "E8EEF5"
LIGHT = "F5F8FB"
WHITE = "FFFFFF"
AMBER = "F4B942"
GREEN = "3A7D44"
RED = "B4433F"


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=85, start=95, bottom=85, end=95) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_twips: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_twips)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "0")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_twips:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_twips[min(index, len(widths_twips) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def keep_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:cantSplit")
    tr_pr.append(marker)


def set_repeat_table_layout(table) -> None:
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def add_field(paragraph, field_code: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = field_code
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, end))


def set_paragraph_border(paragraph, bottom_color: str, size: int = 8) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), bottom_color)
    borders.append(bottom)


def shade_paragraph(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def add_inline_markdown(paragraph, text: str, default_color: str | None = None) -> None:
    text = text.replace("  ", " ").strip()
    token_re = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*)")
    cursor = 0
    for match in token_re.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor : match.start()])
            if default_color:
                run.font.color.rgb = rgb(default_color)
        token = match.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
            run.font.color.rgb = rgb(NAVY)
            run.font.highlight_color = WD_COLOR_INDEX.GRAY_25
        else:
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            if default_color:
                run.font.color.rgb = rgb(default_color)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        if default_color:
            run.font.color.rgb = rgb(default_color)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells):
            rows.append(cells)
        index += 1
    return rows, index


def table_widths(rows: list[list[str]], total: int = 9160) -> list[int]:
    columns = max(len(row) for row in rows)
    weights = []
    for col in range(columns):
        lengths = [len(re.sub(r"[`*]", "", row[col])) if col < len(row) else 0 for row in rows]
        representative = max(8, min(42, int(sum(lengths) / max(1, len(lengths))) + 4))
        weights.append(representative)
    if columns >= 6:
        weights = [max(9, min(24, weight)) for weight in weights]
    weight_sum = sum(weights)
    widths = [max(720, round(total * weight / weight_sum)) for weight in weights]
    delta = total - sum(widths)
    widths[-1] += delta
    return widths


def add_markdown_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    columns = max(len(row) for row in rows)
    if columns > 5:
        headers = rows[0]
        for record_index, source_row in enumerate(rows[1:], start=1):
            label_parts = [re.sub(r"[`*]", "", value) for value in source_row[:2] if value]
            label = " — ".join(label_parts) or f"Record {record_index}"
            label_paragraph = document.add_paragraph()
            label_paragraph.paragraph_format.space_before = Pt(5)
            label_paragraph.paragraph_format.space_after = Pt(2)
            label_paragraph.paragraph_format.keep_with_next = True
            label_run = label_paragraph.add_run(label)
            label_run.bold = True
            label_run.font.size = Pt(8.5)
            label_run.font.color.rgb = rgb(BLUE)

            card = document.add_table(rows=len(headers), cols=2)
            card.style = "Table Grid"
            set_repeat_table_layout(card)
            for field_index, header in enumerate(headers):
                field_cell, value_cell = card.rows[field_index].cells
                value = source_row[field_index] if field_index < len(source_row) else ""
                for cell in (field_cell, value_cell):
                    cell.text = ""
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    set_cell_margins(cell, 60, 85, 60, 85)
                set_cell_shading(field_cell, PALE)
                field_p = field_cell.paragraphs[0]
                value_p = value_cell.paragraphs[0]
                add_inline_markdown(field_p, header)
                add_inline_markdown(value_p, value)
                for run in field_p.runs:
                    run.bold = True
                    run.font.size = Pt(7.4)
                    run.font.color.rgb = rgb(NAVY)
                for run in value_p.runs:
                    run.font.size = Pt(7.4)
                field_p.paragraph_format.space_after = Pt(0)
                value_p.paragraph_format.space_after = Pt(0)
            set_table_geometry(card, [2250, 6910])
        document.add_paragraph().paragraph_format.space_after = Pt(1)
        return
    table = document.add_table(rows=len(rows), cols=columns)
    table.style = "Table Grid"
    set_repeat_table_layout(table)
    font_size = 7.9 if columns <= 4 else 7.0 if columns <= 6 else 6.4
    for row_index, source_row in enumerate(rows):
        target_row = table.rows[row_index]
        for col_index in range(columns):
            cell = target_row.cells[col_index]
            cell.text = ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            add_inline_markdown(paragraph, source_row[col_index] if col_index < len(source_row) else "")
            for run in paragraph.runs:
                run.font.size = Pt(font_size)
                if row_index == 0:
                    run.bold = True
                    run.font.color.rgb = rgb(WHITE)
            if row_index == 0:
                set_cell_shading(cell, NAVY)
            elif row_index % 2 == 0:
                set_cell_shading(cell, LIGHT)
        if row_index == 0:
            repeat_table_header(target_row)
    set_table_geometry(table, table_widths(rows))
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def font(path: str, size: int):
    try:
        return ImageFont.truetype(path, size=size)
    except OSError:
        return ImageFont.load_default()


def clean_node(raw: str) -> str:
    raw = re.sub(r"^[A-Za-z0-9_]+(?=[\[({])", "", raw.strip())
    raw = raw.strip("[](){}\"'")
    raw = raw.replace("<br/>", " / ").replace("_", " ")
    return re.sub(r"\s+", " ", raw).strip() or "Component"


def extract_diagram_edges(block: str) -> list[tuple[str, str, str]]:
    labels: dict[str, str] = {}
    for node_id, label in re.findall(r"\b([A-Za-z][A-Za-z0-9_]*)\s*[\[({]+([^\]\)}]+)[\]\)}]+", block):
        labels[node_id] = clean_node(label)

    edges: list[tuple[str, str, str]] = []
    for raw_line in block.splitlines():
        line = raw_line.strip()
        if not line or line.startswith(("flowchart", "graph", "sequenceDiagram", "stateDiagram", "erDiagram", "subgraph", "end", "participant", "actor", "%%")):
            continue

        sequence = re.match(r"([A-Za-z0-9_]+)\s*-+>>?\s*([A-Za-z0-9_]+)\s*:\s*(.+)", line)
        if sequence:
            source, target, relation = sequence.groups()
            edges.append((labels.get(source, clean_node(source)), labels.get(target, clean_node(target)), relation.strip()))
            continue

        er = re.match(r"([A-Za-z0-9_]+)\s+[|o{}.-]+\s+([A-Za-z0-9_]+)\s*:\s*(.+)", line)
        if er:
            source, target, relation = er.groups()
            edges.append((labels.get(source, clean_node(source)), labels.get(target, clean_node(target)), relation.strip()))
            continue

        state = re.match(r"([^\s]+)\s+--?>\s+([^:]+?)(?:\s*:\s*(.+))?$", line)
        if state:
            source, target, relation = state.groups()
            edges.append((labels.get(source, clean_node(source)), labels.get(target.strip(), clean_node(target)), (relation or "transition").strip()))
            continue

        flow = re.search(r"(.+?)\s+[-.=]+(?:\|([^|]+)\|)?(?:>)\s+(.+)", line)
        if flow:
            source, relation, target = flow.groups()
            source_id = re.match(r"([A-Za-z0-9_]+)", source.strip())
            target_id = re.match(r"([A-Za-z0-9_]+)", target.strip())
            source_name = labels.get(source_id.group(1), clean_node(source)) if source_id else clean_node(source)
            target_name = labels.get(target_id.group(1), clean_node(target)) if target_id else clean_node(target)
            edges.append((source_name, target_name, (relation or "feeds").strip()))

    deduped: list[tuple[str, str, str]] = []
    for edge in edges:
        normalized = tuple(value[:95] for value in edge)
        if normalized not in deduped:
            deduped.append(normalized)
    return deduped[:18]


def wrap_text(draw: ImageDraw.ImageDraw, value: str, text_font, max_width: int, max_lines: int = 2) -> list[str]:
    words = value.split()
    lines: list[str] = []
    current = ""
    for word in words:
        trial = f"{current} {word}".strip()
        if draw.textbbox((0, 0), trial, font=text_font)[2] <= max_width:
            current = trial
        else:
            if current:
                lines.append(current)
            current = word
            if len(lines) == max_lines - 1:
                break
    if current and len(lines) < max_lines:
        lines.append(current)
    remaining = " ".join(words)
    shown = " ".join(lines)
    if len(shown) < len(remaining) and lines:
        lines[-1] = lines[-1].rstrip(".")[: max(3, len(lines[-1]) - 1)] + "..."
    return lines or [value[:30]]


def centered_lines(draw, box, lines, text_font, fill):
    left, top, right, bottom = box
    line_height = text_font.size + 6 if hasattr(text_font, "size") else 24
    y = top + ((bottom - top) - line_height * len(lines)) / 2
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=text_font)
        x = left + ((right - left) - (bbox[2] - bbox[0])) / 2
        draw.text((x, y), line, font=text_font, fill=fill)
        y += line_height


def render_diagram(block: str, title: str, index: int, output_dir: Path) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    edges = extract_diagram_edges(block)
    if not edges:
        statements = [clean_node(line) for line in block.splitlines()[1:] if line.strip() and not line.strip().startswith(("subgraph", "end"))]
        edges = [("Architecture", statement, "contains") for statement in statements[:12]]

    columns = 2 if len(edges) > 5 else 1
    rows = math.ceil(len(edges) / columns)
    width = 1800
    card_width = 820 if columns == 2 else 1660
    card_height = 138
    gap_x = 40
    gap_y = 24
    header_height = 155
    footer_height = 65
    height = header_height + rows * (card_height + gap_y) + footer_height
    image = Image.new("RGB", (width, height), "#F5F8FB")
    draw = ImageDraw.Draw(image)
    regular = font(r"C:\Windows\Fonts\segoeui.ttf", 23)
    small = font(r"C:\Windows\Fonts\segoeui.ttf", 18)
    semibold = font(r"C:\Windows\Fonts\seguisb.ttf", 26)
    title_font = font(r"C:\Windows\Fonts\seguisb.ttf", 39)

    draw.rectangle((0, 0, width, 125), fill=f"#{NAVY}")
    draw.text((58, 30), f"D{index:02d}", font=semibold, fill=f"#{AMBER}")
    title_lines = wrap_text(draw, title, title_font, 1460, 1)
    draw.text((160, 24), title_lines[0], font=title_font, fill="white")

    for item_index, (source, target, relation) in enumerate(edges):
        column = item_index % columns
        row = item_index // columns
        left = 50 + column * (card_width + gap_x)
        top = header_height + row * (card_height + gap_y)
        right = left + card_width
        bottom = top + card_height
        draw.rounded_rectangle((left, top, right, bottom), radius=18, fill="white", outline="#C5D2DE", width=2)
        node_width = int(card_width * 0.36)
        source_box = (left + 18, top + 20, left + 18 + node_width, bottom - 20)
        target_box = (right - 18 - node_width, top + 20, right - 18, bottom - 20)
        draw.rounded_rectangle(source_box, radius=13, fill=f"#{PALE}", outline=f"#{BLUE}", width=2)
        draw.rounded_rectangle(target_box, radius=13, fill="#E5F2F3", outline=f"#{TEAL}", width=2)
        arrow_start = source_box[2] + 15
        arrow_end = target_box[0] - 15
        arrow_y = top + 60
        draw.line((arrow_start, arrow_y, arrow_end, arrow_y), fill=f"#{MUTED}", width=4)
        draw.polygon(((arrow_end, arrow_y), (arrow_end - 14, arrow_y - 9), (arrow_end - 14, arrow_y + 9)), fill=f"#{MUTED}")
        relation_lines = wrap_text(draw, relation, small, max(70, arrow_end - arrow_start - 12), 2)
        rel_y = arrow_y + 12
        for rel_line in relation_lines:
            bbox = draw.textbbox((0, 0), rel_line, font=small)
            draw.text(((arrow_start + arrow_end - (bbox[2] - bbox[0])) / 2, rel_y), rel_line, font=small, fill=f"#{MUTED}")
            rel_y += 21
        centered_lines(draw, source_box, wrap_text(draw, source, regular, node_width - 25), regular, f"#{INK}")
        centered_lines(draw, target_box, wrap_text(draw, target, regular, node_width - 25), regular, f"#{INK}")

    footer_y = height - footer_height + 16
    draw.text((50, footer_y), "Rendered architecture equivalent from the canonical Markdown source", font=small, fill=f"#{MUTED}")
    output = output_dir / f"diagram-{index:02d}.png"
    image.save(output, "PNG", optimize=True)
    return output


def configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(9.6)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_after = Pt(5.5)
    normal.paragraph_format.line_spacing = 1.15

    for style_name, size, color, before, after in (
        ("Title", 30, NAVY, 0, 12),
        ("Subtitle", 14, MUTED, 0, 8),
        ("Heading 1", 17, NAVY, 16, 9),
        ("Heading 2", 13, BLUE, 13, 7),
        ("Heading 3", 10.5, TEAL, 10, 5),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(9.4)
        style.paragraph_format.left_indent = Inches(0.25)
        style.paragraph_format.first_line_indent = Inches(-0.16)
        style.paragraph_format.space_after = Pt(3)

    if "Code Block" not in styles:
        code_style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code_style = styles["Code Block"]
    code_style.font.name = "Consolas"
    code_style.font.size = Pt(7.4)
    code_style.font.color.rgb = rgb(NAVY)
    code_style.paragraph_format.left_indent = Inches(0.12)
    code_style.paragraph_format.right_indent = Inches(0.12)
    code_style.paragraph_format.space_before = Pt(3)
    code_style.paragraph_format.space_after = Pt(3)
    code_style.paragraph_format.keep_together = False


def configure_sections(document: Document) -> None:
    for section in document.sections:
        section.top_margin = Inches(0.72)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.82)
        section.right_margin = Inches(0.82)
        section.header_distance = Inches(0.3)
        section.footer_distance = Inches(0.3)
        section.different_first_page_header_footer = True

        header = section.header
        paragraph = header.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = paragraph.add_run("COSMICFORGE HEALTHNET  /  MARKETING INTELLIGENCE PHASE 2")
        run.font.name = "Calibri"
        run.font.size = Pt(7.5)
        run.font.bold = True
        run.font.color.rgb = rgb(MUTED)
        set_paragraph_border(paragraph, PALE, 6)

        footer = section.footer
        table = footer.add_table(rows=1, cols=2, width=Inches(6.4))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        set_table_geometry(table, [6500, 2660])
        left, right = table.rows[0].cells
        left.text = "Controlled release candidate  |  Internal engineering reference"
        for run in left.paragraphs[0].runs:
            run.font.size = Pt(7.2)
            run.font.color.rgb = rgb(MUTED)
        right_p = right.paragraphs[0]
        right_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        right_p.add_run("Page ")
        add_field(right_p, "PAGE")
        for run in right_p.runs:
            run.font.size = Pt(7.2)
            run.font.color.rgb = rgb(MUTED)


def add_cover(document: Document) -> None:
    kicker = document.add_paragraph()
    kicker.paragraph_format.space_before = Pt(12)
    run = kicker.add_run("COSMICFORGE HEALTHNET CENTRALIZED ADMIN")
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = rgb(TEAL)
    set_paragraph_border(kicker, TEAL, 12)

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(78)

    title = document.add_paragraph(style="Title")
    title.add_run("Marketing Intelligence\nPhase 2")
    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.add_run("Complete Architecture and Implementation")

    descriptor = document.add_paragraph()
    descriptor.paragraph_format.space_before = Pt(18)
    descriptor.paragraph_format.space_after = Pt(42)
    run = descriptor.add_run("Canonical engineering, deployment, live-validation, operations, risk, and acceptance reference")
    run.font.size = Pt(11)
    run.font.color.rgb = rgb(MUTED)

    status = document.add_table(rows=2, cols=2)
    status.style = "Table Grid"
    status_data = (("REVISION", "1.0  /  19 July 2026"), ("RELEASE POSTURE", "Engineering complete  /  Live provider validation pending"))
    for row_index, values in enumerate(status_data):
        for col_index, value in enumerate(values):
            cell = status.rows[row_index].cells[col_index]
            cell.text = value
            set_cell_margins(cell, 130, 150, 130, 150)
            set_cell_shading(cell, NAVY if col_index == 0 else LIGHT)
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(8 if col_index == 0 else 9)
                run.font.bold = col_index == 0
                run.font.color.rgb = rgb(WHITE if col_index == 0 else INK)
    set_table_geometry(status, [2150, 7010])

    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(54)
    note.paragraph_format.keep_together = True
    run = note.add_run("STAGE 1 DARK LAUNCH ONLY")
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = rgb(NAVY)
    note.add_run("\nProvider enablement remains gated by explicit permission, validated credentials, and one-platform-at-a-time evidence.")
    for run in note.runs[1:]:
        run.font.size = Pt(9)
        run.font.color.rgb = rgb(MUTED)
    document.add_page_break()


def add_document_control(document: Document, toc_entries: list[str]) -> None:
    document.add_heading("Document control", level=1)
    intro = document.add_paragraph()
    add_inline_markdown(intro, "Repository: `ADMIN-Backend`  |  Canonical source: `CosmicForge_Marketing_Intelligence_Phase2_Complete_Architecture.md`")
    add_markdown_table(
        document,
        [
            ["Revision", "Date", "Scope", "Status"],
            ["1.0", "2026-07-19", "WP0 through Sections 22-26", "Controlled release candidate"],
        ],
    )

    document.add_heading("Table of contents", level=1)
    instruction = document.add_paragraph()
    run = instruction.add_run("This static contents list is retained for reliable viewing across Word-compatible applications.")
    run.italic = True
    run.font.size = Pt(8.5)
    run.font.color.rgb = rgb(MUTED)
    for entry in toc_entries:
        paragraph = document.add_paragraph(style="List Number")
        paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(6.0))
        add_inline_markdown(paragraph, re.sub(r"^\d+\.\s*", "", entry))
    document.add_page_break()


def build_document(source: Path, output: Path, diagram_dir: Path) -> None:
    raw = source.read_text(encoding="utf-8")
    lines = raw.splitlines()
    toc_entries = [match.group(1) for line in lines if (match := re.match(r"^##\s+(\d+\.\s+.+)$", line))]

    document = Document()
    configure_styles(document)
    section = document.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    add_cover(document)
    add_document_control(document, toc_entries)

    start = next(index for index, line in enumerate(lines) if line.startswith("## 1. Executive overview"))
    index = start
    diagram_index = 0
    current_section = "Architecture"
    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()
        if not stripped:
            index += 1
            continue

        if stripped.startswith("```"):
            language = stripped[3:].strip().lower()
            block_lines: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                block_lines.append(lines[index])
                index += 1
            block = "\n".join(block_lines)
            if language == "mermaid":
                diagram_index += 1
                path = render_diagram(block, current_section, diagram_index, diagram_dir)
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = paragraph.add_run()
                run.add_picture(str(path), width=Inches(6.25))
                caption = document.add_paragraph()
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                caption.paragraph_format.keep_with_next = False
                caption.paragraph_format.space_after = Pt(8)
                caption_run = caption.add_run(f"Figure {diagram_index}. {current_section}")
                caption_run.italic = True
                caption_run.font.size = Pt(8)
                caption_run.font.color.rgb = rgb(MUTED)
            else:
                paragraph = document.add_paragraph(style="Code Block")
                shade_paragraph(paragraph, PALE)
                paragraph.add_run(block)
            index += 1
            continue

        if stripped.startswith("|"):
            rows, index = parse_table(lines, index)
            add_markdown_table(document, rows)
            continue

        heading = re.match(r"^(#{2,4})\s+(.+)$", stripped)
        if heading:
            hashes, heading_text = heading.groups()
            level = len(hashes) - 1
            if level == 1:
                current_section = heading_text
                if document.paragraphs and heading_text != "1. Executive overview":
                    document.add_page_break()
            paragraph = document.add_heading(level=level)
            add_inline_markdown(paragraph, heading_text)
            if level == 1:
                set_paragraph_border(paragraph, PALE, 8)
            index += 1
            continue

        bullet = re.match(r"^-\s+(?:\[([ xX])\]\s+)?(.+)$", stripped)
        if bullet:
            checked, text_value = bullet.groups()
            paragraph = document.add_paragraph(style="List Bullet")
            if checked is not None:
                marker = "☑ " if checked.lower() == "x" else "☐ "
                marker_run = paragraph.add_run(marker)
                marker_run.font.color.rgb = rgb(GREEN if checked.lower() == "x" else MUTED)
            add_inline_markdown(paragraph, text_value)
            index += 1
            continue

        numbered = re.match(r"^\d+\.\s+(.+)$", stripped)
        if numbered:
            paragraph = document.add_paragraph(style="List Number")
            add_inline_markdown(paragraph, numbered.group(1))
            index += 1
            continue

        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if not candidate or candidate.startswith(("#", "|", "```", "- ")) or re.match(r"^\d+\.\s+", candidate):
                break
            paragraph_lines.append(candidate)
            index += 1
        paragraph = document.add_paragraph()
        add_inline_markdown(paragraph, " ".join(paragraph_lines))

    if diagram_index != 14:
        raise RuntimeError(f"Expected 14 Mermaid diagrams, rendered {diagram_index}")

    configure_sections(document)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.core_properties.title = "Marketing Intelligence Phase 2: Complete Architecture and Implementation"
    document.core_properties.subject = "CosmicForge HealthNet Centralized Admin engineering and release reference"
    document.core_properties.author = "CosmicForge HealthNet Engineering"
    document.core_properties.keywords = "marketing intelligence, phase 2, architecture, release, operations"
    document.core_properties.comments = "Generated from the canonical repository Markdown source."
    document.save(output)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--diagram-dir", type=Path, default=DEFAULT_DIAGRAM_DIR)
    args = parser.parse_args()
    build_document(args.source.resolve(), args.output.resolve(), args.diagram_dir.resolve())
    print(args.output.resolve())


if __name__ == "__main__":
    main()
